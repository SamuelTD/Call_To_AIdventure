"""Generate editable PowerPoint, matching PDF and speaker notes. No app dependency changes."""
from pathlib import Path
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = Path(__file__).resolve().parent
BG, PANEL, WHITE, MUTED, ACCENT, GOLD = '101C2B', '1C2C40', 'F4F6FA', 'B5C2D3', '60DFBD', 'F2C879'
slides = []
def add(tag, title, subtitle, cards, takeaway, source, notes, seconds=60):
    slides.append(dict(tag=tag,title=title,subtitle=subtitle,cards=cards,takeaway=takeaway,source=source,notes=notes,seconds=seconds))

add('E1 · BLOC 1', 'Call To AIdventure', 'Des données brutes à un catalogue de monstres exploitable',
    [('COLLECTER', 'Des fichiers sources identifiés et conservés.'), ('FIABILISER', 'Des règles explicites de validation et de fusion.'), ('PARTAGER', 'Un stockage SQLite et une API REST documentée.')],
    'Collecte · SQL · Agrégation · Stockage · API', 'Référentiel Dev IA 2023, p. 1–6 ; règlement, p. 9 et 12',
    'Présenter votre nom et votre rôle. Le fil conducteur est le catalogue de monstres utilisé par un jeu narratif avec IA. Le support adopte le format E1 isolé : 15 minutes, puis 10 minutes de questions. Pour un passage du titre complet, ajuster le minutage avec le centre.', 30)
add('CONTEXTE', 'Une donnée fiable pour un jeu cohérent', 'Le besoin métier guide les choix techniques.',
    [('JOUEUR', 'Affronter des monstres aux statistiques exploitables : points de vie, armure, difficulté.'), ('DÉVELOPPEMENT', 'Disposer d’un catalogue normalisé, réutilisable par le jeu et consultable par API.'), ('EXPLOITATION', 'Reconstruire les données et expliquer leur origine après chaque import.')],
    'L’IA raconte ; les données structurées alimentent les règles du jeu.', 'README.md ; docs/block3/project-brief-requirements.md',
    'Décrire le projet scolaire, les joueurs, le développeur et le jury. Le moteur narratif et le RAG donnent le contexte ; E1 porte ici sur les flux de données des monstres. Ne pas présenter le catalogue SQLite comme la base vectorielle du RAG.', 45)
add('CADRAGE', 'Un périmètre local et reproductible', 'Python · Pydantic · SQLite · Django · OpenAPI',
    [('CONTRAINTES', 'Quelques centaines de monstres. Exécution locale. Démonstration du pipeline sans appel au LLM.'), ('ORGANISATION', 'Projet solo avec assistance outillée. Étapes : collecte, qualité, stockage, exposition, preuves.'), ('À PERSONNALISER', 'Nom et rôle exacts ; dates des étapes ; temps réellement passé ; budget engagé.')],
    'Le budget et le calendrier réels ne sont pas chiffrés dans les documents disponibles.', 'pyproject.toml ; docs/block3/agile-execution.md',
    'Compléter avant soutenance les éléments personnels attendus par le référentiel : commanditaire, contribution, planning et budget. Ne pas inventer de coût nul : distinguer poste local, hébergement éventuel et appels IA du jeu. L’ordre des étapes est un plan de présentation, pas un historique daté.', 45)
add('C1 → C5', 'Une chaîne de données traçable', 'Chaque étape produit une donnée ou une preuve inspectable.',
    [('01 / SOURCES', 'Snapshot collecté + fichier corrigé\n↓\nChargement et métadonnées'), ('02 / QUALITÉ', 'Normalisation + validation\n↓\nFusion, rejets, manifeste'), ('03 / DIFFUSION', 'SQLite + provenance\n↓\nRequêtes SQL et API Django')],
    'Brut conservé → jeu nettoyé → catalogue partagé', 'src/data_pipeline/ ; docs/block1/data-pipeline.md',
    'Suivre le trajet d’un monstre. Les fichiers bruts permettent d’inspecter l’entrée ; les données nettoyées portent le résultat de fusion ; le manifeste résume le traitement ; SQLite et l’API servent les consommateurs.', 60)
add('C1 · EXTRACTION', 'Deux entrées, une origine à expliciter', 'Les snapshots sont reproductibles ; la collecte web reste un sujet distinct.',
    [('SNAPSHOT EXTERNE', 'monster_scrapping/monsters.json\n510 enregistrements\nStatistiques issues de la collecte.'), ('SOURCE CORRIGÉE', 'data/documents/monsters.json\n510 enregistrements\nAdaptations aux besoins du jeu.'), ('CONTRAINTES WEB', 'Scraper documenté : temporisation, timeout, retries et robots. Preuve hors ligne par fixture.')],
    'Les deux fichiers sont liés : leur acceptation comme sources distinctes reste à confirmer.', 'docs/block1/data-pipeline.md ; src/data_pipeline/sources.py',
    'Le run vérifié charge des snapshots locaux : il ne prouve pas un scraping en direct. La documentation signale une page de vérification côté site. Aucune nouvelle collecte réseau n’a été exécutée pour ce support. Distinguer source de collecte et API d’exposition. Les exigences service web et big data ne sont pas couvertes par ces seuls fichiers.', 70)
add('C3 · NORMALISATION', 'Transformer sans perdre le sens', 'Exemple pédagogique repris des cas de test.',
    [('AVANT', 'Nom : Élite Goblin\nHP : "12"\nDifficulté : "1/4"\nDescription : " test "'), ('APRÈS', 'Clé : elite-goblin\nHP : 12 (entier)\nDifficulté : 0,25 + texte "1/4"\nDescription : "test"'), ('CONTRÔLES', 'Nom, armure et HP obligatoires.\nHP strictement positifs.\nErreur → rejet avec motif\nVALIDATION_ERROR.')],
    'Un enregistrement invalide est isolé et conservé pour diagnostic.', 'src/data_pipeline/normalize.py ; tests/test_data_pipeline.py',
    'Expliquer la clé stable sans accent, le typage numérique et la conservation de la fraction initiale. Les valeurs absentes de difficulté peuvent devenir NULL. Le parsing extrait la composante entière des HP/armure : il ne s’agit pas d’une compréhension de toute expression de dés. Le cas invalide est testé même si les snapshots ne génèrent aucun rejet.', 65)
add('C3 · AGRÉGATION', 'Résoudre les conflits de façon déterministe', 'Même clé stable → un seul monstre canonique.',
    [('SOURCE COLLECTÉE', 'Exemple de test : Goblin\nHP = 10\nProvenance conservée.'), ('SOURCE CORRIGÉE', 'Exemple de test : Goblin\nHP = 20\nPriorité aux données corrigées.'), ('RÉSULTAT', 'Goblin : HP = 20\nDeux liens de provenance.\nSource gagnante marquée selected.')],
    'La priorité est une règle métier explicite, pas une preuve de vérité de la source.', 'src/data_pipeline/pipeline.py ; tests/test_data_pipeline.py',
    'Présenter cet exemple comme une fixture de test et non comme un conflit mesuré sur le catalogue réel. La fusion se fait sur la clé de nom normalisée. Limite : cette clé peut rapprocher des homonymes ; une identité métier plus robuste serait nécessaire si le périmètre évolue.', 60)
add('PREUVE · EXÉCUTION', '1 020 entrées deviennent 510 monstres', 'Exécution locale du 14 septembre 2026 · snapshots du dépôt',
    [('1 020', 'Enregistrements acceptés\n1 020 collectés\n0 rejet sur ces snapshots'), ('510', 'Monstres après fusion\n5 conflits détectés\nPriorité à la source corrigée'), ('5 / 5', 'Tests du pipeline réussis\nNormalisation, fractions, fusion,\nprovenance et sorties JSONL')],
    'Le succès du run ne démontre pas, à lui seul, la couverture complète de C1 à C5.', 'preuves/20260914T093833Z-b0d37253/manifest.json ; tests/test_data_pipeline.py',
    'Le run de preuve a utilisé --no-database : il vérifie les sorties sans modifier la base du jeu. Le test de persistance crée une base temporaire et vérifie les sources et rejets. Les cinq tests ne sont pas la suite complète API, RGPD ou application. Aucun temps de performance ni résultat de tests API n’est revendiqué ici.', 55)
add('C2 · REQUÊTES SQL', 'Extraire une donnée utile et explicable', 'Quatre requêtes versionnées couvrent les usages du catalogue.',
    [('Q1 / Q2', 'Q1 : sélectionner un catalogue valide et trié.\nQ2 : joindre chaque monstre à ses sources via un ID paramétré.'), ('Q3 / Q4', 'Q3 : agréger le bilan d’import.\nQ4 : regrouper par difficulté et calculer les moyennes HP / armure.'), ('INDEX', 'Nom insensible à la casse.\nDifficulté.\nLiens de provenance.\nPlans via EXPLAIN QUERY PLAN.')],
    'Les index répondent aux accès ; aucun gain de performance non mesuré n’est annoncé.', 'db/sqlite/queries/monster_dataset.sql ; src/data_pipeline/storage.py',
    'Expliquer SELECT, WHERE, JOIN et GROUP BY avec les finalités métier. Q2 utilise :monster_id. Les index évitent des parcours inutiles selon le plan choisi. La commande explain_queries.py est disponible, mais aucun avant/après de performance n’a été exécuté pour ce support. Une annexe contient la requête exacte.', 70)
add('C4 · MODÉLISATION', 'Séparer le catalogue et les données de compte', 'Vue simplifiée du modèle ; MCD et MLD complets dans la documentation.',
    [('IMPORT', 'INGESTION_RUN\n1 → N MONSTER\n1 → N REJECTED_RECORD\n1 → N MONSTER_SOURCE'), ('PROVENANCE', 'MONSTER\n1 → N MONSTER_SOURCE\n\nClés étrangères, unicité de stable_key, contraintes HP / armure.'), ('DONNÉES PERSONNELLES', 'USER\n1 → N SAVE_GAME\n1 → N CHARACTER_TEMPLATE\n\nBase Django séparée du catalogue.')],
    'Chaque donnée a un rôle, des contraintes et une relation définis.', 'docs/block1/data-models.md ; src/data_pipeline/storage.py',
    'Distinguer MCD (entités et associations), MLD (clés et relations) et modèle physique (DDL SQLite). Chaque source acceptée est liée au monstre canonique et à l’import. Les rejets sont rattachés à l’import. Le modèle affiché est simplifié et ne remplace pas les modèles complets du dossier.', 70)
add('C4 · STOCKAGE', 'SQLite répond au volume de la démonstration', 'Un choix proportionné, avec des limites opérationnelles identifiées.',
    [('POURQUOI SQLITE', 'Fichier local, peu de données, import par un seul processus et lectures fréquentes.'), ('RECONSTRUCTION', 'Sources versionnées + pipeline.\nSchéma créé par le code.\nImport transactionnel avec provenance.'), ('LIMITES', 'Le catalogue est remplacé à chaque import. Historique complet non garanti. PostgreSQL à envisager pour plusieurs écrivains.')],
    'Les données de compte nécessitent une sauvegarde distincte et une restauration testée.', 'src/data_pipeline/storage.py ; docs/block1/backup-restore.md',
    'Le code remplace monsters et monster_sources. Ne pas affirmer qu’un ancien run permet de relire toutes ses anciennes lignées dans SQLite : les artefacts de run sont à conserver. Les IDs numériques du catalogue ne sont pas une identité stable entre imports. La procédure de restauration Django est documentée, pas exécutée pendant la création des slides.', 55)
add('C4 · DONNÉES PERSONNELLES', 'Documenter les traitements et les droits', 'État du registre et des procédures du projet.',
    [('PÉRIMÈTRE', 'Comptes, e-mails, mots de passe hachés, sauvegardes et personnages rattachés à un compte.'), ('PROCÉDURES', 'Export utilisateur.\nRectification via administration.\nSuppression du compte et cascades.\nNettoyage avec prévisualisation.'), ('À VALIDER', 'Bases légales et durées proposées.\nSauvegardes terminées : 365 jours.\nLogs : 30 jours proposés.\nValidation par le responsable.')],
    'Un registre et des commandes constituent des éléments de démarche ; pas une certification de conformité.', 'docs/block1/gdpr-register.md ; docs/block1/backup-restore.md',
    'Restituer le registre interne, sans présenter les durées proposées comme des obligations légales ou une conformité acquise. Montrer la différence entre données de référence et données liées aux comptes. La purge se prévisualise sans --apply. Ne jamais afficher un véritable export utilisateur devant le jury.', 60)
add('C5 · API REST', 'Partager le catalogue avec un contrat clair', 'API en lecture seule · version v1 · documentation OpenAPI',
    [('LECTURE PUBLIQUE', 'GET /api/v1/monsters/\nGET /api/v1/monsters/{id}/\n\nFiltres, tri et pagination\n100 résultats maximum par page.'), ('ACCÈS RESTREINT', 'Résumé d’import :\n/api/v1/ingestion-runs/\n{run_id}/summary/\n\nSession Django + compte staff.'), ('CONTRAT', '/api/v1/docs/\n/api/v1/openapi.yaml\n\nParamètres validés ; SQL paramétré ; erreurs structurées.')],
    'Public pour le catalogue ; réservé au personnel pour les informations d’import.', 'docs/block1/openapi.yaml ; docs/block1/api-security.md',
    'Exemple de lecture : /api/v1/monsters/?page_size=5&challenge_max=1. Expliquer 400 pour paramètre invalide, 404 pour absence, 401/403 pour accès au résumé. Le rate limiting distribué du catalogue est une mesure prévue au reverse proxy, pas une protection démontrée ici. La documentation interactive dépend du CDN, contrairement au contrat YAML local.', 70)
add('DÉMONSTRATION', 'Suivre un résultat de bout en bout', 'Séquence proposée : moins d’une minute, avec preuves hors ligne en secours.',
    [('1 / MANIFESTE', 'Afficher les compteurs :\n1 020 → 510\n5 conflits, 0 rejet.\nOuvrir un objet nettoyé.'), ('2 / PROVENANCE', 'Montrer le test de fusion.\nPuis Q2 : un monstre, ses sources et la source sélectionnée.'), ('3 / API', 'Ouvrir la collection filtrée.\nAfficher le détail d’un monstre.\nRelier les champs au contrat OpenAPI.')],
    'Préparer la base et le serveur avant l’oral ; conserver le manifeste comme solution de secours.', 'Notes orales : commandes et prérequis ; preuves/ ; docs/block1/openapi.yaml',
    'Ne pas lancer une reconstruction --reset sur une base utile pendant l’oral. Préparer un environnement de démonstration séparé. L’API nécessite Django configuré et une base construite ; elle n’a pas été démarrée pour générer ces slides. Si le serveur manque, montrer le contrat et annoncer explicitement que ce n’est pas une réponse HTTP en direct.', 45)
add('BILAN · C1 À C5', 'Un socle concret, des écarts à traiter', 'La couverture technique doit rester distincte de la validation par le jury.',
    [('DÉMONTRABLE', 'Pipeline sur fichiers, normalisation, fusion et persistance testées. SQL, modèles et contrat API présents.'), ('ÉCARTS E1', 'Aucune source big data.\nPas de collecte depuis une API externe.\nDeux snapshots liés.\nCollecte web directe non démontrée.'), ('AVANT L’ORAL', 'Compléter budget et planning.\nRépéter API et restauration.\nValider les choix de sources.\nRapprocher chaque preuve du référentiel.')],
    'Le résultat : un catalogue exploitable dont les transformations sont explicables.', 'Référentiel, C1–C5 ; docs/block1/ ; preuves/manifest.json (run horodaté)',
    'Conclure sur le résultat métier et les choix justifiés. Ne pas affirmer que le faible volume dispense des exigences big data écrites dans le référentiel. Faire traiter cet écart avec le centre. Inviter le jury à examiner les annexes sur SQL, les preuves et les arbitrages.', 40)
add('ANNEXE A · SQL', 'Revenir à l’origine d’un monstre', 'Extrait exact de Q2, catalogue SQL du dépôt.',
    [('SÉLECTION', 'SELECT m.name, s.source,\n s.source_record_id, s.source_url,\n s.collected_at, s.selected,\n s.ingestion_run_id'), ('JOINTURE', 'FROM monsters AS m\nJOIN monster_sources AS s\n ON s.monster_id = m.id'), ('PARAMÈTRE ET TRI', 'WHERE m.id = :monster_id\nORDER BY s.selected DESC,\n s.source;\n\nID lié comme paramètre SQL.')],
    'Le résultat explique à la fois la provenance et la priorité de fusion.', 'db/sqlite/queries/monster_dataset.sql, Q2', 'Afficher cette annexe pour expliquer une jointure, la liaison du paramètre et le tri de la source sélectionnée. Ne pas remplacer le paramètre par une concaténation de texte.', 0)
add('ANNEXE B · PREUVES', 'Où retrouver les éléments du dossier', 'Les chemins sont relatifs à la racine du dépôt.',
    [('C1 / C3', 'src/data_pipeline/\ntests/test_data_pipeline.py\ndocs/block1/data-pipeline.md\npreuves/<run>/manifest.json'), ('C2 / C4', 'db/sqlite/queries/\ndocs/block1/data-models.md\ndocs/block1/database-and-sql.md\ndocs/block1/gdpr-register.md'), ('C5 / CADRAGE', 'docs/block1/openapi.yaml\ndocs/block1/api-security.md\ncertif/*Referentiel*.pdf\ncertif/*Règlement*.pdf')],
    'Preuves générées dans certif/presentation-e1/preuves/.', 'Dépôt local et documents de certification fournis', 'Les documents locaux de certification sont la référence de ce support ; aucune vérification d’une version réglementaire plus récente n’a été réalisée. Le rapport professionnel individuel reste un livrable distinct des slides.', 0)
add('ANNEXE C · QUESTIONS', 'Défendre les choix et leurs limites', 'Quatre questions probables du jury.',
    [('POURQUOI CES SOURCES ?', 'Elles répondent au besoin du jeu. Leur lien est déclaré ; les catégories de sources manquantes restent un écart.'), ('POURQUOI SQLITE ?', 'Volume modeste et démo locale. Plusieurs écrivains ou instances conduiraient à revoir le SGBD.'), ('ET LA QUALITÉ ?', 'Règles déterministes, rejets conservés, provenance et tests. Aucun rejet sur un run ne signifie pas absence de défaut possible.')],
    'Pourquoi l’API est publique ? Elle expose un catalogue de référence ; les résumés d’import sont restreints.', 'docs/block1/ ; tests/test_data_pipeline.py', 'Préparer aussi : collision de clés stables, IDs renouvelés après import, conservation des anciens runs, validation des droits d’exploitation des sources, sécurité et durées des données personnelles.', 0)

slides[13]['seconds'] = 105
slides[13]['subtitle'] = 'Séquence proposée : 1 min 45, avec preuves hors ligne en secours.'
assert sum(s['seconds'] for s in slides)==900
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
pdfmetrics.registerFont(TTFont('Deck', 'C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('DeckBold', 'C:/Windows/Fonts/arialbd.ttf'))
pdf = canvas.Canvas(str(OUT/'Presentation_E1.pdf'), pagesize=(960,540))
pdf.setTitle('Call To AIdventure — Présentation E1')

def rect(slide,x,y,w,h,color):
    shape=slide.shapes.add_shape(1, Inches(x/72), Inches(y/72), Inches(w/72), Inches(h/72))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color); shape.line.fill.background()
    pdf.setFillColor('#'+color); pdf.rect(x,540-y-h,w,h,fill=1,stroke=0)

def txt(slide,x,y,w,h,value,size=20,color=WHITE,bold=False):
    font='DeckBold' if bold else 'Deck'
    lines=[]
    for paragraph in value.split('\n'):
        line=''
        for word in paragraph.split():
            candidate=(line+' '+word).strip()
            if pdfmetrics.stringWidth(candidate,font,size)>w-3 and line:
                lines.append(line); line=word
            else: line=candidate
        lines.append(line)
    if len(lines)*size*1.25>h+1:
        raise ValueError(f'Text overflow: {value}')
    box=slide.shapes.add_textbox(Inches(x/72),Inches(y/72), Inches(w/72), Inches(h/72))
    tf=box.text_frame; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.word_wrap=False
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color); p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=Pt(size*1.25)
    pdf.setFillColor('#'+color); pdf.setFont(font,size)
    for i,line in enumerate(lines): pdf.drawString(x,540-y-size-i*size*1.25,line)

notes=['# Notes orales — Présentation E1', '', '15 diapositives principales : 15 minutes. 3 annexes pour les questions.', '',
       'À personnaliser avant passage : identité, rôle exact, commanditaire, budget réel et dates du projet. Le règlement local prévoit 15 min + 10 min de questions pour le bloc 1 isolé.', '']
for i,s in enumerate(slides):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide,0,0,960,540,BG); rect(slide,36,32,34,4,ACCENT)
    txt(slide,82,22,800,30,s['tag'],12,ACCENT,True)
    txt(slide,36,64,885,90,s['title'],34,WHITE,True)
    txt(slide,36,153,885,55,s['subtitle'],18,MUTED)
    for j,(heading,body) in enumerate(s['cards']):
        x=36+j*300
        rect(slide,x,224,288,205,PANEL)
        txt(slide,x+16,240,254,44,heading,18,ACCENT,True)
        txt(slide,x+16,290,254,130,body,16,WHITE)
    txt(slide,36,450,888,52,s['takeaway'],17,GOLD,True)
    txt(slide,36,515,850,18,s['source'],9,MUTED)
    txt(slide,904,510,42,24,f'{i+1:02d}',14,ACCENT,True)
    slide.notes_slide.notes_text_frame.text=s['notes']+'\n\nSources : '+s['source']
    pdf.showPage()
    notes.extend([f"## {i+1:02d}. {s['title']}"+(f" — {s['seconds']} s" if s['seconds'] else ' — annexe'),'',s['notes'],'',f"Sources : {s['source']}",''])
prs.save(OUT/'Presentation_E1.pptx'); pdf.save()
notes.extend(['## Démonstration — commandes PowerShell', '', 'Depuis la racine du dépôt, pour régénérer les artefacts sans modifier la base du jeu :', '', '```powershell', "$env:PYTHONPATH='src'", 'uv run python -m data_pipeline --no-database --output-dir certif/presentation-e1/preuves', 'uv run python -m unittest discover -s tests -p test_data_pipeline.py', '```', '', 'API : préparer au préalable Django et une base de démonstration selon README.md. Ouvrir `/api/v1/monsters/?page_size=5&challenge_max=1`, puis le détail d’un ID renvoyé et `/api/v1/docs/`. Ne pas supposer qu’un ID est conservé après reconstruction.', '', '## Vérification réalisée', '', '14 septembre 2026 : pipeline sur snapshots locaux, sans écriture dans la base du jeu. 1 020 collectés et acceptés, 0 rejet, 510 fusionnés, 5 conflits. Suite `tests/test_data_pipeline.py` : 5 tests réussis. API et restauration non exécutées pour ce support.', '', '## Régénérer le support', '', '```powershell', 'uv run --with python-pptx --with reportlab python certif/presentation-e1/generate_slides.py', '```', '', 'Le PDF reprend la même mise en page que le PowerPoint. Tous les textes et blocs du PowerPoint sont modifiables. Les notes orales figurent également dans les notes du présentateur.'])
(OUT/'Notes_orales_E1.md').write_text('\n'.join(notes),encoding='utf-8')
print(f'Generated {len(slides)} slides; main presentation: 900 seconds.')
