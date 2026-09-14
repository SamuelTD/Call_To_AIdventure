"""Presentations for E2 and E3; uses the existing E1 visual primitives only."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT=Path(__file__).resolve().parent
BG,PANEL,WHITE,MUTED,ACCENT,GOLD='101C2B','1C2C40','F4F6FA','B5C2D3','60DFBD','F2C879'
# Reuse drawing helpers without executing or rewriting the E1 presentation.
source=(OUT.parent/'presentation-e1/generate_slides.py').read_text(encoding='utf-8-sig')
exec(source[source.index('def rect('):source.index("notes=['# Notes orales")])
pdfmetrics.registerFont(TTFont('Deck','C:/Windows/Fonts/arial.ttf'))
pdfmetrics.registerFont(TTFont('DeckBold','C:/Windows/Fonts/arialbd.ttf'))
slides=[]
def add(part,tag,title,subtitle,cards,takeaway,source,notes):
    slides.append(dict(part=part,tag=tag,title=title,subtitle=subtitle,cards=cards,takeaway=takeaway,source=source,notes=notes,seconds=90 if part=='E2' else 100 if part=='E3' else 0))

add('E2','E2 · C6 À C8','Choisir et configurer le service d’IA','Call To AIdventure · Bloc 2 · E2 : 15 minutes',
 [('BESOIN','Un maître du jeu narratif en français et en anglais, relié à un état de jeu.'),('DÉMARCHE','Veille, critères de sélection, architecture et configuration.'),('POSITIONNEMENT','Intégration de services préexistants. Pas d’entraînement d’un modèle de génération dans ce projet.')],
 'De la demande fonctionnelle à une solution intégrable.', 'Référentiel 2023, p. 6–9 ; règlement, p. 12',
 'Présenter votre identité, votre rôle exact et le commanditaire. E2 dure 15 minutes ; E3 suit pendant 20 minutes, puis le bloc prévoit 10 minutes de questions. Le rapport professionnel reste distinct de ce support. Le choix du service doit être défendu par des preuves, pas uniquement par son utilisation dans le code.')
add('E2','C7 · EXPRESSION DU BESOIN','Raconter, proposer, agir : trois contrats','L’IA intervient dans un moteur dont les règles restent côté serveur.',
 [('NARRATION','État + action + contexte documentaire → suite de l’aventure en FR ou EN.'),('CHOIX','Exactement 3 actions distinctes.\n6 mots maximum par action.\nStructure validée par Pydantic.'),('DÉCISIONS','Objectifs et salle : sorties typées.\nCombat, soin et dégâts : outils autorisés et règles déterministes.')],
 'Une sortie grammaticalement correcte peut rester incorrecte pour le jeu.', 'docs/block2/ai-requirements-and-selection.md ; src/agents/schemas.py',
 'Illustrer avec une porte fermée : la narration décrit la situation, le modèle propose trois actions, puis le joueur en choisit une. Le modèle ne doit pas inventer directement les nouvelles statistiques. Les schémas imposent une forme, pas la pertinence narrative ou la langue : celles-ci nécessitent une évaluation distincte.')
add('E2','C7 · CRITÈRES','Définir la réussite avant de comparer','Les valeurs suivantes sont des objectifs documentés, pas des scores mesurés.',
 [('QUALITÉ','100 % de sorties structurées valides.\nAu moins 90 % de réussite sur le jeu de cas.\nRelecture humaine FR / EN.'),('SERVICE','Objectif p95 inférieur à 30 s.\nDisponibilité logique visée : 99 %.\nÉchec contrôlé sans perte d’état.'),('CONTRAINTES','Budget mensuel à renseigner.\nConditions de traitement à vérifier.\nMatériel local pour les embeddings.')],
 'Séparer validité du format, qualité de la réponse et disponibilité du service.', 'docs/block2/ai-requirements-and-selection.md ; ai-test-strategy.md',
 'Expliquer p95 : 95 % des observations sont sous cette durée sur une fenêtre définie. Ne pas confondre timeout d’un appel et durée totale d’un tour, qui peut contenir plusieurs appels et reprises. Budget, volume réel et arbitrage final ne sont pas fournis. Le taux de 90 % exige un corpus représentatif et une notation adaptée.')
add('E2','C6 · VEILLE','Une veille à rattacher aux décisions','Trame proposée ; les preuves de veille ne figurent pas dans le dépôt consulté.',
 [('THÈMES','Sorties structurées et outils.\nRAG et évaluation.\nSécurité, données personnelles, accessibilité et sobriété.'),('MÉTHODE PROPOSÉE','1 h hebdomadaire.\nAgrégateur RSS + journal daté.\nAuteur, date et fiabilité.\nRecoupement des informations.'),('PREUVE ATTENDUE','Une synthèse accessible.\nUne recommandation argumentée.\nUne décision liée au projet.\nUne trace de partage effective.')],
 'À compléter avec votre veille réelle ; aucune synthèse diffusée n’est inventée.', 'Référentiel, C6, p. 6–8 ; certif/block-2-technical-gap-analysis.md',
 'Le référentiel demande au minimum une heure hebdomadaire. Le document de diagnostic renvoie C6 à un dossier séparé : son absence dans ce dépôt ne prouve pas que vous ne l’avez pas réalisé ailleurs. Insérer vos dates, outils, sources identifiées, synthèses et destinataires. Cette slide est une trame, pas une veille effectivement conduite ni un état de l’art actualisé.')
add('E2','C7 · COMPARAISON','Comparer sur un protocole identique','Le benchmark réel de plusieurs modèles reste à réaliser.',
 [('MODÈLE HÉBERGÉ ACTUEL','Intégré dans le code.\nStructure, langue, latence et coût à mesurer.\nChoix encore provisoire.'),('SECOND SERVICE HÉBERGÉ','Candidat exact à nommer.\nMême dataset et même révision.\nComparer les capacités requises et les conditions de traitement.'),('GÉNÉRATION LOCALE','Option à étudier.\nMatériel et latence à mesurer.\nDonnées locales ; exploitation à assurer.')],
 'Aucun service n’est déclaré meilleur sans résultats comparables.', 'docs/block2/ai-requirements-and-selection.md',
 'Pour chaque candidat, relever version, date, validité, qualité FR/EN, p50/p95, tokens, coût estimé, région de traitement et rétention. Inclure les informations disponibles sur la démarche écoresponsable. Expliciter les services non étudiés et pourquoi. Ce support ne remplace pas le benchmark exigé par C7 ; aucune invocation payante n’a été réalisée.')
add('E2','C7 · ARCHITECTURE','Une génération hébergée, un RAG local','Décision d’architecture déjà documentée dans le projet.',
 [('GÉNÉRER','Service OpenAI-compatible\n↓\nNarration, choix et décisions\n\nModèle piloté par configuration.'),('RETROUVER','JSON de lore → fragments\n↓\nEmbeddings Ollama → Chroma\n↓\nContexte filtré par aventure.'),('ORCHESTRER','LangGraph\n↓\nTransitions explicites\n↓\nOutils et moteur de jeu.')],
 'Localiser les embeddings ne signifie pas que le contexte reste entièrement local.', 'docs/block2/ai-requirements-and-selection.md ; docs/rag-system.md',
 'La justification actuelle vise la génération multilingue et structurée sans héberger un gros modèle, un petit corpus vectoriel local et des transitions explicites. Le lore retrouvé est injecté dans les prompts envoyés au service de génération : rendre ce flux clair. Ne pas présenter cette décision qualitative comme un benchmark déjà conclu.')
add('E2','C8 · PARAMÉTRAGE','Rendre la configuration explicite','Valeurs par défaut du code ; l’environnement peut les remplacer.',
 [('GÉNÉRATION','Modèle : OPENAI_MODEL.\nClé dans l’environnement.\nTimeout d’appel : 30 s.\nMétadonnées sans clé secrète.'),('RETRIEVAL','RAG_ENABLED configurable.\nOllama : hôte local par défaut.\nEmbeddings : mxbai-embed-large.\nTimeout : 10 s ; 3 tentatives.'),('FRONTIÈRE API','Choix : 1 000 caractères max.\nQuota par défaut : 20 tours par utilisateur et par heure.\nValidation de configuration.')],
 'Les paramètres décrivent le déploiement ; les secrets restent hors des supports.', 'src/agents/runtime_config.py ; .env.example',
 'Lire les valeurs réellement configurées avant démonstration. Le nom du modèle par défaut est une chaîne du dépôt, pas une vérification de disponibilité commerciale. La commande check_ai_services valide les paramètres ; son option --connect vérifie Ollama, pas une génération réelle du fournisseur. Ne jamais afficher le fichier .env.')
add('E2','C8 · RAG','Préparer le contexte avant la génération','Vérification locale du corpus : 57 fragments préparés.',
 [('CORPUS','6 fichiers de personnages.\n3 fichiers de lieux.\n57 fragments au dry-run.\nMétadonnées de source.'),('À L’EXÉCUTION','Requête contextualisée.\nRecherche filtrée par aventure.\nPassages injectés au prompt.\nCache selon le service.'),('EN CAS D’ÉCHEC','Contexte vide et poursuite du jeu.\nÉchec visible dans les métriques.\nQualité narrative à surveiller.')],
 'Le dry-run vérifie la préparation ; il ne mesure ni les embeddings ni le rappel du RAG.', 'src/retrieval/ ; vérification du 14/09/2026',
 'Le dry-run exécuté prépare 36 fragments de personnages et 21 de lieux. La documentation ancienne annonce 39 : la mesure du code courant prime. Aucun serveur Ollama ni Chroma n’a été testé en ligne ici. L’objectif de rappel des sources doit être mesuré séparément ; le runner actuel ne le calcule pas.')
add('E2','C7 / C8 · ARBITRAGES','Maîtriser coûts, données et dépendances','Des arbitrages à chiffrer avant une sélection définitive.',
 [('COÛT','Tours × appels par tour × tokens moyens × tarifs.\nAjouter hébergement, stockage et calcul local.'),('DONNÉES','Décrire ce qui entre dans les prompts. Vérifier les conditions du fournisseur et les règles de conservation.'),('SOBRIÉTÉ','Limiter le contexte utile.\nRéutiliser les embeddings.\nComparer qualité et ressources.\nBilan environnemental non mesuré.')],
 'Le registre des coûts et l’analyse fournisseur restent à compléter avec des données datées.', 'docs/block2/ai-requirements-and-selection.md ; docs/rag-system.md',
 'La formule complète utilise des tarifs par million de tokens : diviser la consommation par un million et distinguer entrée/sortie. Aucun tarif actuel n’est avancé. Les coûts estimés ne sont pas une facture. Éviter d’affirmer qu’un modèle local est automatiquement moins cher ou plus sobre ; cela dépend du matériel et des usages.')
add('E2','E2 · BILAN','Une solution configurée, une sélection à étayer','Passer de l’architecture à la preuve de pertinence.',
 [('C6','Relier le dossier de veille.\nAjouter synthèses datées et décisions réellement prises.'),('C7','Besoins et protocole présents.\nComparer au moins deux candidats.\nCompléter budget, données et critères environnementaux.'),('C8','Configuration et RAG préparés.\nProuver la connectivité réelle et la génération dans un environnement de démonstration.')],
 'Transition E3 : exposer, intégrer, tester et surveiller le composant.', 'docs/block2/evidence-checklist.md',
 'Rappeler ce qui fonctionne dans le code et ce qui reste à montrer en réel. Préparer une démonstration safe de configuration et une sortie structurée issue d’un vrai appel, en conservant modèle, date et contexte d’évaluation. Le bilan ne prétend pas valider toutes les compétences E2.')

add('E3','E3 · C9 À C13','Mettre le composant d’IA en service','Call To AIdventure · Bloc 2 · E3 : 20 minutes',
 [('EXPOSER / INTÉGRER','API authentifiée, connexion au moteur du jeu et gestion des erreurs.'),('SURVEILLER / TESTER','Métriques techniques, qualité des sorties et vérifications reproductibles.'),('LIVRER','Workflow de validation, image Docker et procédure de promotion / retour arrière.')],
 'Le composant livré comprend le code, les prompts, la configuration et les preuves.', 'Référentiel, C9–C13, p. 10–16 ; docs/block2/traceability.md',
 'Introduire E3 de façon autonome si ce fichier est présenté séparément. Il s’agit de mettre en service un modèle préexistant dans une application, pas d’entraîner ses poids. Expliquer ce qui est versionné localement et ce qui dépend du fournisseur distant.')
add('E3','C9 · API REST','Exposer trois points d’entrée documentés','Contrat versionné dans docs/block2/openapi.yaml.',
 [('ÉTAT DE CONFIGURATION','GET /api/v1/ai/health/\nPublic.\nIndique la configuration.\nNe prouve pas un appel réussi au fournisseur.'),('MÉTADONNÉES','GET /api/v1/ai/configuration/\nAuthentifié.\nModèles et limites utiles.\nAucune clé ni prompt.'),('TOUR DE JEU','POST /api/v1/games/\n{game_id}/turns/\nSession + CSRF.\nAvance une sauvegarde possédée.')],
 'Une configuration valide ne constitue pas une mesure de disponibilité du modèle.', 'src/django/game/ai_api.py ; docs/block2/openapi.yaml',
 'Le health renvoie des booléens, pas un ping fournisseur. Décrire le body {"choice": "Open the gate"}. L’API réutilise StepGameView pour déclencher le moteur. Le contrat YAML existe, mais quelques écarts contrat/code sont identifiés en annexe et ne doivent pas être présentés comme une conformité OpenAPI parfaite.')
add('E3','C9 · CONTRÔLES D’ACCÈS','Contrôler avant de solliciter le modèle','La frontière v1 associe identité, propriété et limites.',
 [('IDENTITÉ / PROPRIÉTÉ','Session Django.\nRecherche par ID + utilisateur.\nPartie étrangère ou absente : même réponse 404.'),('REQUÊTE','CSRF requis pour le POST.\nTaille et longueur bornées.\nChoix comparé à la liste serveur lorsqu’elle est non vide.'),('QUOTA','Compteur utilisateur en cache.\nDépassement : HTTP 429.\nUn cache partagé et sa concurrence sont à revoir pour plusieurs instances.')],
 'Les tests prouvent des contrôles précis ; ils ne constituent pas un audit complet de sécurité.', 'src/django/game/ai_api.py ; src/django/game/test_ai_api.py',
 'Résultats vérifiés : utilisateur étranger rejeté, choix arbitraire rejeté avec liste renseignée, CSRF refusé sans jeton et quota exercé. Limite importante : le code ne rejette pas par ce contrôle un choix lorsque current_choices est vide. Ne pas affirmer une prévention générale des injections de prompt. Ne pas montrer d’identifiants de session réels.')
add('E3','C10 · INTÉGRATION','Relier l’interface au moteur narratif','Deux surfaces HTTP réutilisent la logique de jeu.',
 [('NAVIGATEUR','play.html → api_step\n↓\nStepGameView\n\nParcours actif du jeu.'),('API DE CERTIFICATION','/api/v1/games/{id}/turns/\n↓\nContrôles puis StepGameView\n\nSurface v1 distincte.'),('MOTEUR COMMUN','GameEngine → LangGraph\n↓\nRAG, modèle et outils\n↓\nRécit / combat / fin / erreur.')],
 'Le navigateur n’appelle pas directement l’endpoint v1 dans le code consulté.', 'src/django/game/templates/game/play.html ; ai_api.py ; services/game_engine.py',
 'Éviter un schéma trompeur qui ferait passer le navigateur par v1 : le template pointe vers api_step. L’intégration effective du fournisseur passe par les chaînes du moteur ; v1 expose une autre frontière. Une preuve C10 complète doit montrer les endpoints réellement consommés et les adaptations d’interface, avec leur évaluation d’accessibilité.')
add('E3','C10 · SORTIES ET ERREURS','Valider la forme et protéger la progression','Les appels au modèle peuvent échouer ou produire une réponse inutilisable.',
 [('SORTIES TYPÉES','ChoiceOutput : 3 choix distincts, chacun ≤ 6 mots.\nObjectifs et progression : schémas dédiés.'),('REPRISES','Retry sur erreurs transitoires.\nDélais et tentatives bornés.\nAprès épuisement : erreur contrôlée du service.'),('ÉTAT DU JEU','Une génération échouée ne doit pas avancer l’état.\nLes règles de soin, dégâts et progression restent côté serveur.')],
 'Réussite HTTP, validité du schéma et pertinence narrative sont trois vérifications différentes.', 'src/agents/schemas.py ; src/django/game/tests.py',
 'Les tests ciblés ont exercé les prompts FR, la reprise narrative, l’échec sans progression, les outils de soin et les objectifs. La classe SaveGamePersistenceTests n’appartient pas au lot de 33 exécuté ici : ne pas attribuer à ce lot la totalité des tests de persistance. Les cas sont simulés ; aucune qualité linguistique en réel n’est mesurée.')
add('E3','C11 · OBSERVABILITÉ','Observer disponibilité, qualité et coût','Collecteurs → Prometheus → Grafana ; règles vers Alertmanager.',
 [('SERVICE','Nombre d’appels et erreurs.\nHistogramme des durées.\nObjectifs : 99 % de succès et p95 < 30 s.'),('QUALITÉ / RAG','Sorties structurées valides ou invalides.\nRAG : résultats, vides, erreurs et durées de récupération.'),('CONSOMMATION','Tokens d’entrée / sortie.\nCoût estimé avec tarifs configurés.\nLabels sans prompts, identités ni sessions.')],
 'Les SLO sont des objectifs ; aucun historique de production n’est présenté ici.', 'src/observability/ ; docs/block2/monitoring-slos.md ; monitoring/',
 'Expliquer compteur et histogramme, ainsi que la fenêtre de calcul. Les métriques exposées ont des tests ciblés, mais la chaîne Prometheus/Grafana n’a pas été lancée pour ce support. Les tokens peuvent dépendre des données d’usage fournies par le service ; l’estimation de coût nécessite des tarifs à jour. Les labels bornés évitent données sensibles et cardinalité excessive.')
add('E3','C11 · ALERTES','Transformer un signal en action','Règles présentes ; livraison d’une notification externe à démontrer.',
 [('DÉTECTER','3 indisponibilités en 10 min.\nLatence élevée persistante.\nSortie structurée invalide.\nErreurs de retrieval.'),('DIAGNOSTIQUER','Identifier modèle et révision.\nComparer erreurs, durée et RAG.\nReproduire avec un cas minimal sans donnée personnelle.'),('AMÉLIORER','Corriger prompt ou configuration.\nRejouer les cas de référence.\nComparer avant / après.\nGarder un retour arrière.')],
 'Alertmanager utilise un récepteur local : aucune notification externe n’a été prouvée.', 'docs/block2/monitoring-slos.md ; monitoring/alertmanager/',
 'Le seuil précis de RAG documenté est plus de deux échecs en quinze minutes. La latence doit être élevée pendant dix minutes. Préparer un déclenchement contrôlé en bac à sable avec capture horodatée et retour à la règle initiale. Ce projet ajuste l’intégration ; il ne réentraîne pas automatiquement le modèle distant.')
add('E3','C12 · STRATÉGIE DE TEST','Distinguer tests logiciels et évaluation IA','Trois niveaux complémentaires, avec des limites explicites.',
 [('INTÉGRATION SIMULÉE','Tests Django et mocks.\nAuthentification, CSRF, quota, erreurs, prompts, outils et transitions.'),('DATASET HORS LIGNE','4 cas déclarés : choix EN, choix FR, source RAG et entrée adverse.\nValidation du format du dataset.'),('ÉVALUATION RÉELLE','--live appelle les 2 cas de choix.\nLes cas RAG et frontière API restent marqués hors ligne.\nQualité sémantique non notée.')],
 'Le runner actuel ne mesure pas un taux global de qualité du modèle.', 'src/evaluation/runner.py ; evaluation/dataset.json',
 'Point central : validate_dataset vérifie version, identifiants et types de tâches ; ce n’est pas une validation exhaustive du contenu. En live, passed=True signifie ici passage de la chaîne et de ChoiceOutput. Les seuils 1,0 et 0,9 sont écrits dans le rapport ; le code de sortie exige tous les résultats passed. Ne pas annoncer 100 % de qualité à partir de validated_offline.')
add('E3','C12 · PREUVES LOCALES','33 tests ciblés réussis','Vérifications du 14 septembre 2026 · appels IA simulés',
 [('33 TESTS','API et configuration.\nPrompts FR, scope / cache RAG.\nSoin, objectifs et métriques.\nDjango : aucun problème détecté.'),('57 FRAGMENTS','Préparation du corpus réussie.\nAucun calcul d’embedding pendant le dry-run.\nAucune mesure de rappel.'),('4 CAS DÉCLARÉS','Rapport offline-validation.\nPas de benchmark fournisseur.\nPas de score FR / EN.\nPas de latence modèle mesurée.')],
 'Ces résultats démontrent un socle logiciel ; ils ne remplacent pas l’évaluation en conditions réelles.', 'preuves/evaluation-offline.json ; preuves/verification.md',
 'Nommer les classes exécutées, disponibles dans le fichier de preuve. Les 33 tests sont un sous-ensemble, pas toute la suite ni une mesure de couverture. Le résultat de 3,711 s est la durée des tests, jamais une latence d’inférence. Aucune suite API complète ou audit de sécurité ne doit être revendiqué à partir de ce lot.')
add('E3','C13 · INTÉGRATION CONTINUE','Automatiser les contrôles avant livraison','Workflow Block 2 AI component versionné dans GitHub Actions.',
 [('QUALITÉ','Dépendances verrouillées.\nLint et checks Django.\nTests + seuil de couverture 55 %.\nRAG dry-run + dataset offline.'),('ARTEFACT','Validation des configurations de monitoring.\nBuild Docker tagué SHA Git.\nSmoke test HTTP /health.'),('ÉVALUATION LIVE','Déclenchement manuel.\nEnvironnement ai-evaluation.\nClé protégée.\nRapport conservé comme artefact.')],
 'Le workflow est présent ; aucun succès GitHub distant n’a été vérifié pour ces slides.', '.github/workflows/block2-ai.yml ; docs/block2/delivery.md',
 'Décrire PR, push main et déclenchement manuel. Le conteneur est construit et chargé localement dans le job : le workflow ne prouve pas la publication d’une image dans un registre. Le smoke test utilise une clé factice et RAG désactivé ; il vérifie le démarrage HTTP, pas le modèle. La couverture 55 % est un seuil configuré, pas une mesure obtenue aujourd’hui.')
add('E3','C13 · LIVRAISON','Promouvoir un artefact, prévoir le retour arrière','Procédure documentée ; déploiement réel encore à prouver.',
 [('PROMOTION','Revoir les tests et le rapport IA.\nPublier l’image exacte.\nDéployer son digest en staging.\nVérifier santé et parcours.'),('TRAÇABILITÉ','Révision Git et dataset.\nModèle et paramètres.\nPrompts versionnés.\nCollection RAG identifiée.\nConserver les rapports.'),('RETOUR ARRIÈRE','Redéployer le digest précédent.\nRestaurer sa configuration.\nRevenir à la collection RAG prévue.\nRejouer santé et smoke tests.')],
 'Un digest fixe le conteneur ; il ne fige pas à lui seul le comportement d’un fournisseur distant.', 'docs/block2/delivery.md ; src/evaluation/runner.py',
 'Garder au moins un artefact connu bon. Le modèle hébergé peut évoluer si son identifiant n’est pas une version figée ; expliciter cette limite. Versionner une collection RAG fait partie de la procédure à organiser, pas d’une preuve d’activation déjà produite ici. Staging, registre et liens de jobs restent à renseigner.')
add('E3','E3 · DÉMONSTRATION ET BILAN','Montrer une preuve à chaque frontière','Séquence proposée sur un environnement de démonstration préparé.',
 [('API / INTÉGRATION','Montrer configuration sans secret.\nExercer un tour autorisé.\nAfficher un refus de propriété.\nExpliquer les routes navigateur / v1.'),('TESTS / SURVEILLANCE','Afficher le lot de tests et le rapport offline.\nEn environnement prêt : métriques, erreur simulée et alerte.'),('LIVRAISON','Ouvrir le workflow.\nRelier code et artefact.\nPrésenter un déploiement et un rollback seulement s’ils ont été exécutés.')],
 'Restent à prouver : qualité réelle, accessibilité, alerte reçue et livraison effective.', 'docs/block2/evidence-checklist.md ; Notes_orales_Bloc2.md',
 'Les manipulations tiennent dans le créneau de cette slide si les onglets sont prêts. Sans service démarré, montrer les preuves hors ligne et nommer explicitement ce qui est une simulation. Ne pas déclencher un appel payant ou envoyer une notification à un tiers par surprise. Préparer l’audit d’accessibilité et les liens de démonstration avant la soutenance.')

add('ANNEXE','ANNEXE A · C9','Écarts repérés entre contrat et code','Points identifiés à la lecture ; aucune modification applicative dans cette livraison.',
 [('OPENAPI','La liste security encode session OU CSRF, alors que la vue exige les deux.\nRéunir les exigences dans le contrat.'),('ENTRÉES','Le contrat interdit les champs supplémentaires ; la vue ne les rejette pas explicitement.\nUn JSON non objet reste à tester.'),('CHOIX','La comparaison est conditionnée à une liste serveur non vide.\nDéfinir et tester le comportement quand elle est vide.')],
 'Conserver ces réserves dans l’argumentaire de sécurité.', 'docs/block2/openapi.yaml ; src/django/game/ai_api.py',
 'Ces observations proviennent du code local, sans test d’exploitation. Le lot courant couvre une liste de choix non vide. Ajouter des tests de contrat et corriger les écarts avant de revendiquer une conformité stricte. Ce support ne change pas le périmètre de la demande en tâche de développement.')
add('ANNEXE','ANNEXE B · TRAÇABILITÉ','Retrouver les preuves C6 à C13','Sources relatives à la racine du dépôt.',
 [('E2','C6 : dossier de veille à joindre.\nC7 : ai-requirements-and-selection.md\nC8 : runtime_config.py et documentation RAG.'),('E3 · SERVICE','C9 : ai_api.py + test_ai_api.py\nC10 : game_engine.py et play.html\nC11 : observability/ et monitoring/'),('E3 · QUALITÉ','C12 : evaluation/ et tests Django\nC13 : block2-ai.yml\nRapport : preuves/evaluation-offline.json')],
 'La documentation de synthèse se trouve sous docs/block2/.', 'Référentiel fourni, p. 6–16 ; docs/block2/traceability.md',
 'Les chemins des cartes sont abrégés pour projection. Utiliser les notes et le dépôt pour ouvrir les fichiers exacts. Le référentiel local fourni est la base de ce support ; aucune affirmation n’est faite sur une éventuelle version plus récente de la certification.')
add('ANNEXE','ANNEXE C · QUESTIONS','Défendre les limites du dispositif','Préparer une réponse courte, puis une preuve.',
 [('POURQUOI CE MODÈLE ?','Architecture cohérente avec le besoin ; sélection finale à soutenir par un comparatif réel.'),('COMMENT MESURER ?','Format par schéma ; pertinence par cas et revue humaine ; disponibilité et délai par métriques.'),('QU’EST-CE QUI MANQUE ?','Veille jointe, benchmark, notation RAG / adversarial, audit accessible, alerte reçue et déploiement tracé.')],
 'Un modèle préentraîné déplace les contrôles vers les données, les prompts et l’intégration.', 'docs/block2/evidence-checklist.md ; src/evaluation/runner.py',
 'Préparer aussi : confidentialité du contexte envoyé, timeout par appel versus tour complet, quota multi-instance, limite des schémas et stabilité des modèles hébergés. Rester précis sur l’absence d’entraînement ou de réentraînement dans ce projet.')

def render(name, selected):
    global pdf
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    pdf=canvas.Canvas(str(OUT/f'{name}.pdf'),pagesize=(960,540)); pdf.setTitle(name.replace('_',' '))
    notes=[f'# Notes orales — {name}', '', 'Personnaliser identité, commanditaire, rôle, dates, budget et preuves externes avant passage.', '']
    for i,s in enumerate(selected):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); rect(slide,0,0,960,540,BG); rect(slide,36,32,34,4,ACCENT)
        txt(slide,82,22,800,30,s['tag'],12,ACCENT,True)
        txt(slide,36,64,885,90,s['title'],34,WHITE,True)
        txt(slide,36,153,885,55,s['subtitle'],18,MUTED)
        for j,(heading,body) in enumerate(s['cards']):
            x=36+j*300; rect(slide,x,224,288,205,PANEL)
            txt(slide,x+16,240,254,44,heading,18,ACCENT,True)
            txt(slide,x+16,290,254,130,body,16,WHITE)
        txt(slide,36,450,888,52,s['takeaway'],17,GOLD,True)
        txt(slide,36,515,850,18,s['source'],9,MUTED)
        txt(slide,904,510,42,24,f'{i+1:02d}',14,ACCENT,True)
        slide.notes_slide.notes_text_frame.text=f"Durée indicative : {s['seconds']} s.\n\n{s['notes']}\n\nSources : {s['source']}"
        pdf.showPage()
        notes.extend([f"## {i+1:02d}. {s['title']} — {s['seconds']} s",'',s['notes'],'',f"Sources : {s['source']}",''])
    prs.save(OUT/f'{name}.pptx'); pdf.save()
    (OUT/f'Notes_{name}.md').write_text('\n'.join(notes),encoding='utf-8')
    print(f'{name}: {len(selected)} slides, {sum(s["seconds"] for s in selected)} seconds')

assert sum(s['seconds'] for s in slides if s['part']=='E2')==900
assert sum(s['seconds'] for s in slides if s['part']=='E3')==1200
render('Presentation_Bloc2',slides)
render('Presentation_E2',[s for s in slides if s['part']=='E2'])
render('Presentation_E3',[s for s in slides if s['part'] in {'E3','ANNEXE'}])
(OUT/'Notes_orales_Bloc2.md').write_text((OUT/'Notes_Presentation_Bloc2.md').read_text(encoding='utf-8'),encoding='utf-8')
